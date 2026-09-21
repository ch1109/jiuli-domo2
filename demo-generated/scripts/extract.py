"""只读抽取；所有输出均在 demo-generated，原件绝不写入。"""
from pathlib import Path
import sys,json,hashlib,subprocess
from concurrent.futures import ThreadPoolExecutor
ROOT=Path(__file__).resolve().parents[2]; OUT=ROOT/'demo-generated'
sys.path.insert(0,str(OUT/'.python-deps'))
from pypdf import PdfReader
from openpyxl import load_workbook
import xlrd

def dump(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str))
def run(p):
 rel=p.relative_to(ROOT).as_posix(); fid='F-'+hashlib.sha256(rel.encode()).hexdigest()[:12]
 dest=OUT/'extracted'/f'{fid}.json'
 if dest.exists():
  cached=json.loads(dest.read_text())
  if cached['sha256']!=hashlib.sha256(p.read_bytes()).hexdigest():raise RuntimeError('源文件与已复核抽取不一致，必须重新人工复核：'+rel)
  return cached
 obj={'id':fid,'path':rel,'sampleId':p.relative_to(ROOT/'真实整单样本').parts[0], 'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'size':p.stat().st_size,'extension':p.suffix.lower()}
 if '核对单' in p.parts or 'output' in p.parts: obj['role']='reference'
 elif '查货文件' in p.parts or (p.parent.name=='input' and p.suffix.lower()=='.pdf'):obj['role']='inspection'
 else:obj['role']='entrustment_material'
 if p.suffix.lower()=='.xlsx':
  w=load_workbook(p,data_only=False); cached=load_workbook(p,data_only=True)
  obj['sheets']=[{'name':s.title,'rows':s.max_row,'columns':s.max_column,'merged':[str(x) for x in s.merged_cells.ranges], 'cells':[{'cell':c.coordinate,'row':c.row,'column':c.column,'value':c.value,'cached':cached[s.title][c.coordinate].value} for row in s for c in row if c.value is not None]} for s in w]
 elif p.suffix.lower()=='.xls':
  w=xlrd.open_workbook(p,formatting_info=True)
  obj['sheets']=[{'name':s.name,'rows':s.nrows,'columns':s.ncols,'merged':s.merged_cells,'cells':[{'row':i+1,'column':j+1,'value':s.cell_value(i,j)} for i in range(s.nrows) for j in range(s.ncols) if s.cell_value(i,j)!='']} for s in w.sheets()]
 elif p.suffix.lower() in ['.pdf','.jpg','.png']:
  pages=[]; renders=OUT/'renders'/fid; renders.mkdir(exist_ok=True)
  if p.suffix.lower()=='.pdf':
   r=PdfReader(p)
   subprocess.run(['pdftoppm','-scale-to','2100','-png',str(p),str(renders/'page')],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
   imgs=sorted(renders.glob('page-*.png'),key=lambda q:int(q.stem.split('-')[-1]))
   texts=[q.extract_text() or '' for q in r.pages]
  else:imgs=[p];texts=['']
  result=subprocess.run([str(OUT/'scripts'/'.ocr-local'),*[str(q) for q in imgs]],capture_output=True,text=True,check=True)
  obs=[json.loads(x) for x in result.stdout.splitlines()]
  for i,(img,t,o) in enumerate(zip(imgs,texts,obs),1): pages.append({'page':i,'renderPath':str(img.relative_to(ROOT)),'text':t,'ocr':o.get('observations',[]),'error':o.get('error')})
  obj['pages']=pages
 dump(dest,obj)
 print(fid,obj['sampleId'],p.name,flush=True)
 return obj
if __name__=='__main__':
 files=sorted(p for p in (ROOT/'真实整单样本').rglob('*') if p.is_file() and p.suffix.lower() in ['.xlsx','.xls','.pdf','.jpg','.png'])
 with ThreadPoolExecutor(max_workers=3) as ex: result=list(ex.map(run,files))
 dump(OUT/'extracted'/'index.json',[{k:x[k] for k in ['id','path','sampleId','role','sha256','extension']} for x in result])
 print('DONE',len(result))
