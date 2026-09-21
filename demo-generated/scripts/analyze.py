"""阶段一：先建立样本盘点和逐行事实，不生成Mock。"""
from reviewed_facts import *
from collections import Counter,defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from build_mock import template_columns
FIELDS=template_columns()
NUM={'数量','报关单价','总价','件数','净重','毛重','托盘数','期票天数'}
CUSTOMER_BY_SAMPLE={'2025YBT010-2':YBT,'2026(DG)ZW001':ZW,'2026(DG)ZW003':ZW,'2026(DG)ZW050':ZW,'2026ACSY003':'深圳市澳创实业有限公司','2026AG001':'深圳市傲冠软件股份有限公司','2026BMH001':'深圳市百闽海科技有限公司','2026CNKJ001':'福建省超年科技股份有限公司','26SHPYD056':PY,'英卡-抽+整':YK,'多对多样例':OL}
NOTES={
'2025YBT010-2':['委托型号带#9/#D后缀，查货型号无该后缀；不能将去后缀定义为通用匹配规则。','表格毛重合计9.20，入仓章9，两个量值保留。'],
'2026(DG)ZW001':['委托导单没有客户字段，草稿待补客户；查货Bill/Sold To明确东莞智微。','查货两份箱单均只有共箱重量；不将参考结果中的分摊重量反写到查货行。','UPI印刷COD与手写COO不同证据分别保留。'],
'2026(DG)ZW003':['委托导单没有客户字段，草稿待补客户。','委托两条CS5511AN必须保留；查货允许逻辑合并但共箱重量不可分摊。','查货ASL产地CHINA，委托中国台湾；RTS5411S-GR的COO/COD及手写数量另有差异。'],
'2026(DG)ZW050':['委托导单没有客户字段，草稿待补客户。','查货列为C.O.D. KOREA，不能静默当COO；查货毛重1.56、入仓章2、委托2分别保留。'],
'2026ACSY003':['查货Delivery Note只显示TO:S59，没有客户实名，不凭目录建立客户映射。','委托4行，净毛及箱数按两行共用；参考结果的逐行分摊不作为源事实。','委托MICRON产地SINGAPORE，查货手写COD MALAYSIA，不混淆COO与COD。'],
'2026AG001':['查货License编号2026AG001但未显示客户，等待明确归属；不凭文件夹名入池。','两条同型号委托不合并；服务年限委托一年、查货三年，存在真实冲突。'],
'2026BMH001':['11个查货PDF；印刷买方FUJIAN CENTERM或FUJIAN STAR-NET AIOT，与百闽海不同，查货客户待确认。','Arrow文件6个PDF页，原件印刷PAGE 1~6 OF 8，缺7、8页；122行总数仅计现有页。','Arrow现有印刷数量1800000，委托SN74AHC1G08DBVR为360000；不能按委托改写查货。','同一LM2576原始行数量4000含DE2500/GB1500；不自动按产地拆份。'],
'2026CNKJ001':['委托方福建省超年科技股份有限公司，查货Bill to深圳市快极科技有限公司，归属待确认。','入仓条码26057448和正文018230127属于哪个编号体系UNKNOWN，保留后者为歧义证据，不重复生成商品。','委托品牌广濑、产地韩国；查货品牌HIROSE、产地JAPAN；是否品牌别名待确认。'],
'26SHPYD056':['委托9行、合计166140；查货9条批次明细、合计156140；真实差异10000保留。','委托WINBOND的8行与查货7个批次不是逐行一一对应，不能静默合并委托。','重量主要在箱/整单层级；参考结果合并成2行，不作为委托行来源。'],
'英卡-抽+整':['三份独立委托，每份3行；同名发票箱单在三个目录重复出现，仍记录各次材料关联。','两份查货PDF实际均为26070093、12行UMW2631，不能从文件名制造26070092。','客户由内部单明确关联26070093到英卡佐证；英华国际仅记录材料上客户称谓，不自动建别名。','26070092原始查货缺失；发票箱单8000812735提供40条装箱行，属于辅助材料而非已到查货。','原始26070093共177000，三份UMW2631委托共186000，额外9000见辅助箱单但缺查货入仓证据。','8-14委托UMW2631净重23.9大于毛重20.4；不修正原件。'],
'多对多样例':['委托表只有“附件查看”，无商品表、超链接或内嵌附件；原始委托数及行数UNKNOWN。','内部单21行和4行仅作结果佐证，不转为原始委托。','12个PDF、12个入仓号；两个内部单分别使用不同入仓号集合，不能仅凭目录名宣称自然多对多。','原始COO TH与内部结果美国有冲突；保留产地说明书，不以COD替代COO。']}
def decimal(v):
 if v in [None,'']:return U
 if isinstance(v,(int,float,Decimal)):
  # XLS二进制浮点的尾数保留在rawValue；语义值用15位有效数去除存储噪声，不做业务取整。
  return format(Decimal(format(v,'.15g')) if isinstance(v,float) else Decimal(v),'f')
 return str(v).strip()
def mapping(sample):
 if 'ZW' in sample:return {1:'商品类型',2:'品名',3:'品牌',4:'型号',5:'商品描述',6:'产地',7:'单位',8:'数量',9:'报关单价',10:'件数',11:'净重',12:'毛重',14:'sku',15:'供应商',16:'期票天数',17:'对应的采购',18:'物料号码',19:'托盘数',20:'原箱号',21:'备注'}
 if sample in ['2025YBT010-2','2026BMH001']:return {2:'品名',3:'品牌',4:'型号',5:'商品描述',6:'产地',7:'单位',8:'数量',9:'报关单价',10:'总价',11:'件数',12:'净重',13:'毛重'}
 if sample=='2026AG001':return {2:'品牌',4:'品名',5:'型号',6:'商品描述',7:'产地',8:'数量',9:'报关单价',10:'总价',11:'净重',12:'毛重',13:'件数'}
 if sample=='26SHPYD056':return {2:'品牌',3:'品名',4:'型号',5:'商品描述',6:'数量',7:'报关单价',8:'总价',9:'产地',10:'净重',11:'毛重',12:'件数',13:'入仓号'}
 if sample=='英卡-抽+整':return {2:'供应商',3:'品名',4:'品牌',5:'采购订单号',6:'原YK订单号',7:'物料号码',8:'型号',9:'商品描述',10:'产地',11:'单位',12:'数量',14:'报关单价',15:'总价',16:'件数',17:'净重',18:'毛重'}
 return {}
def extract_drafts():
 drafts=[]
 for d in sorted(DOCS.values(),key=lambda x:x['path']):
  if d['role']!='entrustment_material' or 'sheets' not in d or d['sampleId']=='多对多样例':continue
  sample=d['sampleId']; s=d['sheets'][0]; cells={(c['row'],c['column']):c for c in s['cells']}; mp=mapping(sample)
  modelcol=next(c for c,f in mp.items() if f=='型号'); header=1 if 'ZW' in sample else 6
  did='D-'+d['id'][2:]; cust=U if 'ZW' in sample else CUSTOMER_BY_SAMPLE[sample]
  draft={'id':did,'sampleId':sample,'fileId':d['id'],'customerName':cust,'customerEvidence':{'fileId':d['id'],'sheet':s['name'],'location':'委托方' if cust!=U else '表内无客户字段'},'rows':[],'sharedMeasurements':[]}
  for ri in range(header+1,s['rows']+1):
   if not cells.get((ri,modelcol),{}).get('value'):continue
   if not ('ZW' in sample or isinstance(cells.get((ri,1),{}).get('value'),(int,float))):continue
   vals={f:U for f in FIELDS};evidence={};extra={}
   for ci,field in mp.items():
    cell=cells.get((ri,ci),{});raw=cell.get('value');v=cell.get('cached') if isinstance(raw,str) and raw.startswith('=') else raw
    if field in FIELDS:vals[field]=decimal(v)
    elif v not in [None,'']:extra[field]=decimal(v)
    if v not in [None,'']:evidence[field]={'fileId':d['id'],'sheet':s['name'],'cell':f'{get_column_letter(ci)}{ri}','rawValue':raw,'cachedValue':cell.get('cached'),'value':decimal(v)}
   vals['客户名']=cust
   if cust!=U:evidence['客户名']=draft['customerEvidence']|{'value':cust,'rawValue':cust}
   if 'ZW' not in sample:vals['币种']='USD';evidence['币种']={'fileId':d['id'],'sheet':s['name'],'location':'单价/总价表头USD','value':'USD','rawValue':'USD'}
   # 跨多商品行的合并单元格值属于共享事实，不能归给首行。
   merged=s['merged']
   for area in merged:
    if isinstance(area,str):
     from openpyxl.utils.cell import range_boundaries
     c0,r0,c1,r1=range_boundaries(area)
    else:r0,r1,c0,c1=area;r0+=1;c0+=1
    if r1>r0 and r0<=ri<=r1:
     for ci,field in mp.items():
      if c0<=ci<=c1 and field in ['件数','净重','毛重']:
       anchor=cells.get((r0,c0),{}).get('value')
       if anchor not in [None,'']:draft['sharedMeasurements'].append({'field':field,'value':decimal(anchor),'rows':[r0,r1],'source':{'fileId':d['id'],'sheet':s['name'],'cell':f'{get_column_letter(c0)}{r0}'}})
       vals[field]=U;evidence.pop(field,None)
   draft['rows'].append({'id':f'{did}-R{ri:03}','sourceOrder':len(draft['rows'])+1,'sourceRow':ri,'fields':vals,'evidence':evidence,'extraFields':extra})
  drafts.append(draft)
 # 两份非Excel委托逐行录入，原件有可靠文字或目视读取。
 for fid,cust,items in [
 ('F-0971c3fd7c49',CUSTOMER_BY_SAMPLE['2026ACSY003'],[
 {'品牌':'MICRON','品名':'集成电路','型号':'MT29F4T08EUHAFM4-NW920','产地':'SINGAPORE','数量':'1401','单位':'PCS','报关单价':'57.67','总价':'80795.67'},
 {'品牌':'MICRON','品名':'集成电路','型号':'MT29F2T08CUHBBM4-NW821','产地':'SINGAPORE','数量':'1564','单位':'PCS','报关单价':'28.83','总价':'45090.12'},
 {'品牌':'Toshiba','品名':'集成电路','型号':'TH58TF1V23BA8H-1V23','产地':'JAPAN','数量':'1080','单位':'PCS','报关单价':'14.42','总价':'15573.60'},
 {'品牌':'Toshiba','品名':'集成电路','型号':'TH58TF2T23BA8J-2T23','产地':'JAPAN','数量':'368','单位':'PCS','报关单价':'28.83','总价':'10609.44'}]),
 ('F-a64834ca9065',CUSTOMER_BY_SAMPLE['2026CNKJ001'],[{'品牌':'广濑','品名':'接插件','型号':'BM25U-4P/2-V(85)','产地':'韩国','数量':'580000','报关单价':'0.0397','总价':'23026.00'}])]:
  d=DOCS[fid];did='D-'+fid[2:]; dr={'id':did,'sampleId':d['sampleId'],'fileId':fid,'customerName':cust,'customerEvidence':{'fileId':fid,'page':1,'location':'委托方/公章'},'rows':[],'sharedMeasurements':[]}
  for i,item in enumerate(items,1):
   item=dict(item)
   if d['sampleId']=='2026ACSY003':item['商品描述']='非多元件集成电路；已切割、已封装；应用场景：消费电子；存储功能；容量：'+['512GB','256GB','128GB','256GB'][i-1]+'，量产；非加密；非易失'
   vals={f:U for f in FIELDS};vals.update(item);vals.update({'客户名':cust,'币种':'USD'})
   evidence={f:{'fileId':fid,'page':1,'position':f'委托商品表第{i}行/{f}' if f not in ['客户名','币种'] else '委托方或表头','value':v,'rawValue':v} for f,v in vals.items() if v!=U}
   dr['rows'].append({'id':f'{did}-R{i:03}','sourceOrder':i,'sourceRow':i,'fields':vals,'evidence':evidence,'extraFields':{}})
  if d['sampleId']=='2026ACSY003':dr['sharedMeasurements']=[{'source':{'fileId':fid,'page':1},'rows':[1,2],'净重':'7.0','毛重':'7.6','件数':'1'},{'source':{'fileId':fid,'page':1},'rows':[3,4],'净重':'4.0','毛重':'4.6','件数':'1'}]
  drafts.append(dr)
 return drafts

def build_analysis():
 drafts=extract_drafts();manifest=[];files=[]; seen={}
 for d in sorted(DOCS.values(),key=lambda d:d['path']):
  dup=seen.get(d['sha256']);seen.setdefault(d['sha256'],d['id'])
  files.append({k:d[k] for k in ['id','path','sampleId','extension','sha256','size','role']}|{'pageCount':len(d['pages']) if 'pages' in d else None,'worksheets':[s['name'] for s in d.get('sheets',[])],'byteDuplicateOf':dup,'contentDuplicateOf':'F-a5368260946b' if d['id']=='F-baf3de39ef68' else None,'contentDuplicateEvidence':'条码、发票号、12行和印章均相同，已逐页复核' if d['id']=='F-baf3de39ef68' else None})
 # 已逐页比较的PDF/XLS版本：单证相同，不新增委托行。
 from build_mock import segments
 versions={'F-1f89bab4fe64':'F-1df4f4d83480','F-e3577cb3f191':'F-a235e97d6dd0','F-3c457d2e01db':'F-5a09ab721f2e'}
 for f in files:
  f['documentSegments']=segments(f)
  f['sameBusinessDocumentAs']=versions.get(f['id'])
  if f['id'] in versions:f['versionEvidence']='PDF与XLS商品表均3行；型号、数量、委托日期逐项相符；原始签章PDF独立保留'
 for sample in sorted(set(d['sampleId'] for d in DOCS.values())):
  ds=[d for d in drafts if d['sampleId']==sample];os=[o for o in ORDERS if o['sampleId']==sample];fs=[f for f in files if f['sampleId']==sample]
  merge=any(len([r for r in o['rows'] if (r['品牌'],r['型号'],r['产地'])==(a['品牌'],a['型号'],a['产地'])])>1 and U not in (a['品牌'],a['型号'],a['产地']) for o in os for a in o['rows'])
  m={'sampleId':sample,'customer':CUSTOMER_BY_SAMPLE[sample],'customerDetermination':'来自委托方、查货明确Bill To或内部单；未证明的归属在对象级标UNKNOWN','fileIds':[f['id'] for f in fs], 'entrustmentDocumentCount':len(ds) if ds else U,'invoiceDocumentCount':2 if sample=='英卡-抽+整' else (10 if sample=='多对多样例' else 15 if sample=='2026BMH001' else 0),'packingListDocumentCount':2 if sample=='英卡-抽+整' else 14 if sample=='多对多样例' else 1 if sample=='2026BMH001' else 0 if sample in ['2026AG001','2026ACSY003'] else len(os)+(1 if sample=='26SHPYD056' else 0),'inspectionPdfCount':sum(f['extension']=='.pdf' and f['role']=='inspection' for f in fs),'warehouseCountConfirmed':len(os),'warehouseAmbiguities':['018230127'] if sample=='2026CNKJ001' else [],'logicalInspectionOrderCount':len(os),'entrustmentLineCount':sum(len(d['rows']) for d in ds) if ds else U,'inspectionSourceLineCount':sum(len(o['rows']) for o in os),'hasMergeableGoods':merge,'hasFieldConflict':True if sample in ['2025YBT010-2','2026(DG)ZW003','2026(DG)ZW050','2026AG001','2026BMH001','2026CNKJ001','26SHPYD056','英卡-抽+整','多对多样例'] else U,'hasMultipleCandidatePotential':True if merge else U,'possibleRelationScenarios':(['多对一候选（待客户/逐行关系核实）'] if len(os)>1 else ['一对一候选（待逐行关系核实）'])+(['一对多（有内部单依据，原始查货不全）'] if sample=='英卡-抽+整' else []),'notes':NOTES[sample],'inspectionFacts':os,'entrustmentFacts':ds,'reviewRecords':[REVIEW[f['id']] for f in fs]}
  if sample=='多对多样例':m['possibleRelationScenarios']=['参考结果支持两组多对一；真实多对多UNKNOWN'];m['hasFieldConflict']=True
  manifest.append(m);(OUT/'manifests'/f'{sample}.json').write_text(json.dumps(m,ensure_ascii=False,indent=2))
 inventory={'phase':'Phase -1｜真实样本 → Mock 数据基线','countingPolicy':'文件数按物理文件；入仓号/原始行排除重复库存；委托按独立单据保留原始行；UNKNOWN不是零。发票/箱单包含查货PDF内嵌单证，同一文件可有多个类型。','files':files,'samples':[{k:v for k,v in m.items() if k not in ['inspectionFacts','entrustmentFacts','reviewRecords']} for m in manifest],'summary':{'sampleCount':len(manifest),'materialFileCount':len(files),'pdfCount':sum(f['extension']=='.pdf' for f in files),'pdfPages':sum(f['pageCount'] or 0 for f in files if f['extension']=='.pdf'),'inspectionPdfCount':sum(m['inspectionPdfCount'] for m in manifest),'logicalInspectionOrderCount':len(ORDERS),'inspectionSourceLineCount':sum(len(o['rows']) for o in ORDERS),'entrustmentDraftCountConfirmed':len(drafts),'entrustmentLineCountConfirmed':sum(len(d['rows']) for d in drafts),'entrustmentLineCountUnknownSamples':['多对多样例']}}
 (OUT/'source_inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2))
 lines=['# 真实样本分析','', '阶段：Phase -1｜真实样本 → Mock 数据基线','', '已完整读取72份材料（含50个PDF、94页、20个工作簿及2张图片）。源目录另有21个系统隐藏文件，仅纳入完整性校验。真实样本没有写入操作。','', '## 统计口径','', '- UNKNOWN 表示无法可靠确认，不参与数值求和，不等于0。','- 文件名不决定入仓号；包装/批次明细不与同批发票或送货合计重复入池。','- COO、COD、参考结果和仓库实测值分别保存；没有批准的字段优先级不自动覆盖。','- “客户”列是该样本中已证明的业务客户；对象级UNKNOWN仍须人工补充，不代表整套材料归属已确认。','- 查货原始行按最细印刷商品/批次明细计数；同页已明确重复的LOT附表只作补充证据。','', '|样本|客户|委托份数|发票份数|箱单份数|查货PDF|已确认入仓号/逻辑单|原始查货行|委托行|可合并|','|---|---|---:|---:|---:|---:|---:|---:|---:|---|']
 for m in manifest:lines.append('|'+ '|'.join(str(m[k]) for k in ['sampleId','customer','entrustmentDocumentCount','invoiceDocumentCount','packingListDocumentCount','inspectionPdfCount','logicalInspectionOrderCount','inspectionSourceLineCount','entrustmentLineCount','hasMergeableGoods'])+'|')
 for m in manifest:lines += ['',f"## {m['sampleId']}",'']+['- '+n for n in m['notes']]+[f"- 完整文件、行结构及页码见 manifests/{m['sampleId']}.json。"]
 lines += ['','## 先盘点、后生成','', '本报告与source_inventory在Mock生成前建立。后续build_mock仅消费这些逐行事实，不根据验收结果反改源事实。','', '## 读取限制','', 'OCR只作为定位辅助，关键商品结构经过页面复核。模糊字段和共享重量没有采用模型猜测。原件缺页、缺附件、客户对应关系和手写修改优先级均进入OPEN_QUESTIONS。']
 (OUT/'SAMPLE_ANALYSIS.md').write_text('\n'.join(lines)+'\n')
 print(json.dumps(inventory['summary'],ensure_ascii=False))
if __name__=='__main__':build_analysis()
