#!/usr/bin/env python3
"""只读的文本 PDF / Excel 解析适配器。

它只输出原始事实和来源位置，不执行客户归属、商品匹配或重量推断。
扫描 PDF 没有可提取文本时返回 OCR_REQUIRED，而不是把图片识别结果伪装成已解析事实。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "demo-generated/.python-deps"))

from openpyxl import load_workbook  # type: ignore  # noqa: E402
from pypdf import PdfReader  # type: ignore  # noqa: E402

try:
    import xlrd  # type: ignore  # noqa: E402
except ImportError:
    xlrd = None

ADAPTER_ID = "local-text-material"
ADAPTER_VERSION = "0.1.0"

HEADER_ALIASES = {
    "客户名": {"客户名", "客户", "收货人"}, "品牌": {"品牌", "牌号"},
    "型号": {"型号", "规格型号", "料号", "货号"}, "商品描述": {"商品描述", "描述"},
    "品名": {"品名", "商品名称", "名称"}, "产地": {"产地", "原产地"},
    "单位": {"单位", "计量单位"}, "数量": {"数量", "申报数量", "查货数量"},
    "报关单价": {"报关单价", "单价", "申报单价"}, "总价": {"总价", "金额", "申报总价"},
    "币种": {"币种", "货币"}, "件数": {"件数", "包装件数"}, "净重": {"净重"}, "毛重": {"毛重"},
    "sku": {"sku"}, "对应的采购": {"对应的采购"}, "供应商号码": {"供应商号码"}, "供应商": {"供应商"},
    "期票天数": {"期票天数"}, "采购订单号": {"采购订单号", "订单号"}, "物料号码": {"物料号码"},
    "托盘数": {"托盘数"}, "入仓号": {"入仓号"}, "产线": {"产线", "生产线"}, "备注": {"备注", "说明"},
}

def normalize_header(value: str) -> str:
    return "".join(str(value).strip().split()).replace("（", "").replace("）", "").replace("(", "").replace(")", "").lower()

def detect_headers(values: list[Any]) -> dict[int, str]:
    headers: dict[int, str] = {}
    for index, value in enumerate(values, 1):
        normalized = normalize_header(value or "")
        for field, aliases in HEADER_ALIASES.items():
            if normalized in {normalize_header(alias) for alias in aliases}:
                headers[index] = field
                break
    return headers

def named_fields(values: list[Any], headers: dict[int, str]) -> dict[str, str]:
    fields = {f"column_{index}": value for index, value in enumerate(values, 1) if value not in (None, "")}
    for index, field in headers.items():
        value = values[index - 1] if index - 1 < len(values) else None
        if value not in (None, ""):
            fields[field] = str(value)
    return fields


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def excel_result(path: Path, source_file_id: str, content_hash: str, material_type: str | None, customer_id: str | None) -> dict[str, Any]:
    facts: list[dict[str, Any]] = []
    warnings: list[str] = []
    sheet_metadata: list[dict[str, Any]] = []
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheets = workbook.worksheets
        for sheet in sheets:
            sheet_metadata.append({"name": sheet.title, "merged": [str(item) for item in sheet.merged_cells.ranges]})
            rows = list(sheet.iter_rows(values_only=True))
            headers: dict[int, str] = {}
            header_row = None
            for candidate_number, candidate in enumerate(rows, 1):
                detected = detect_headers(list(candidate))
                if len(detected) >= 2:
                    headers, header_row = detected, candidate_number
                    break
            if header_row is not None:
                header_values = [None if value is None else str(value) for value in rows[header_row - 1]]
                facts.append({
                    "id": f"{source_file_id}-{sheet.title}-HEADER-R{header_row}",
                    "fields": named_fields(header_values, headers),
                    "sourceLocation": {"page": None, "sheet": sheet.title, "row": header_row, "position": None},
                    "confidence": None,
                    "warnings": ["表头行，仅用于保留表格结构，不作为商品事实"],
                })
            for row_number, row in enumerate(rows, 1):
                if row_number == header_row:
                    continue
                values = [None if value is None else str(value) for value in row]
                if not any(value not in (None, "") for value in values):
                    continue
                fields = named_fields(values, headers)
                facts.append({
                    "id": f"{source_file_id}-{sheet.title}-R{row_number}",
                    "fields": fields,
                    "sourceLocation": {"page": None, "sheet": sheet.title, "row": row_number, "position": None},
                    "confidence": None,
                    "warnings": [],
                })
        workbook.close()
    elif suffix == ".xls":
        if xlrd is None:
            raise RuntimeError("XLS_PARSER_UNAVAILABLE")
        workbook = xlrd.open_workbook(path, formatting_info=False)
        for sheet in workbook.sheets():
            sheet_metadata.append({"name": sheet.name, "merged": [list(item) for item in sheet.merged_cells]})
            rows = [sheet.row_values(row_number) for row_number in range(sheet.nrows)]
            headers: dict[int, str] = {}
            header_row = None
            for candidate_number, candidate in enumerate(rows, 1):
                detected = detect_headers(candidate)
                if len(detected) >= 2:
                    headers, header_row = detected, candidate_number
                    break
            if header_row is not None:
                header_values = [None if value == "" else str(value) for value in rows[header_row - 1]]
                facts.append({
                    "id": f"{source_file_id}-{sheet.name}-HEADER-R{header_row}",
                    "fields": named_fields(header_values, headers),
                    "sourceLocation": {"page": None, "sheet": sheet.name, "row": header_row, "position": None},
                    "confidence": None,
                    "warnings": ["表头行，仅用于保留表格结构，不作为商品事实"],
                })
            for row_number, raw_values in enumerate(rows, 1):
                if row_number == header_row:
                    continue
                values = [None if value == "" else str(value) for value in raw_values]
                if not any(value not in (None, "") for value in values):
                    continue
                fields = named_fields(values, headers)
                facts.append({
                    "id": f"{source_file_id}-{sheet.name}-R{row_number + 1}",
                    "fields": fields,
                    "sourceLocation": {"page": None, "sheet": sheet.name, "row": row_number + 1, "position": None},
                    "confidence": None,
                    "warnings": [],
                })
    else:
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    if not facts:
        warnings.append("NO_TABULAR_FACTS")
    return result(source_file_id, content_hash, material_type, customer_id, facts, warnings, {"sheets": sheet_metadata})


def pdf_result(path: Path, source_file_id: str, content_hash: str, material_type: str | None, customer_id: str | None) -> dict[str, Any]:
    facts: list[dict[str, Any]] = []
    warnings: list[str] = []
    reader = PdfReader(path)
    for page_number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        facts.append({
            "id": f"{source_file_id}-P{page_number}",
            "fields": {"rawText": text},
            "sourceLocation": {"page": page_number, "sheet": None, "row": None, "position": None},
            "confidence": None,
            "warnings": ["文本 PDF 原始页事实，未做商品字段推断"],
        })
    if not facts:
        warnings.append("OCR_REQUIRED")
    return result(source_file_id, content_hash, material_type, customer_id, facts, warnings, {"pages": len(reader.pages)})


def result(source_file_id: str, content_hash: str, material_type: str | None, customer_id: str | None, facts: list[dict[str, Any]], warnings: list[str], metadata: dict[str, Any]) -> dict[str, Any]:
    return {
        "adapterId": ADAPTER_ID,
        "adapterVersion": ADAPTER_VERSION,
        "sourceFileId": source_file_id,
        "contentSha256": content_hash,
        "materialType": material_type,
        "customerId": customer_id,
        "customerResolution": "已识别" if customer_id else "待补客户信息",
        "facts": facts,
        "warnings": warnings,
        "metadata": metadata,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("file", type=Path)
    parser.add_argument("--source-file-id", required=True)
    parser.add_argument("--material-type", default=None)
    parser.add_argument("--customer-id", default=None)
    args = parser.parse_args()
    if not args.file.is_file():
        print(json.dumps({"errorCode": "FILE_NOT_FOUND", "errorMessage": str(args.file)}, ensure_ascii=False))
        return 2
    try:
        content_hash = sha256(args.file)
        if args.file.suffix.lower() == ".pdf":
            output = pdf_result(args.file, args.source_file_id, content_hash, args.material_type, args.customer_id)
        elif args.file.suffix.lower() in {".xlsx", ".xls"}:
            output = excel_result(args.file, args.source_file_id, content_hash, args.material_type, args.customer_id)
        else:
            print(json.dumps({"errorCode": "UNSUPPORTED_FILE_TYPE", "errorMessage": args.file.suffix}, ensure_ascii=False))
            return 2
    except Exception as error:  # noqa: BLE001 - CLI boundary must return a stable error contract.
        print(json.dumps({"errorCode": "ADAPTER_FAILED", "errorMessage": str(error)}, ensure_ascii=False))
        return 1
    print(json.dumps(output, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
